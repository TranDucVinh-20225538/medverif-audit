#!/usr/bin/env python3
"""Build the blind IAA v1 annotation kit (annotation/iaa_v1/).

Stratified sample (~150 items) over the main 110 DOSE bases plus the
polarity / rephrase / frequency-route sets.  Each stratum mixes items the
author ACCEPTED and items the author REJECTED (where rejected items exist),
so that Cohen's kappa against the author is meaningful.

Outputs (all under annotation/iaa_v1/):
  annotation_sheet.csv / annotation_sheet.xlsx   -- BLIND, give to annotator
  private_key.jsonl                               -- author labels; DO NOT SHARE
  sample_manifest.json                            -- counts, seed, sources (no labels)

Blinding: no author labels, no rule outputs, no instance ids / base ids in
the sheet; items shuffled; opaque ids (IAA1-xxxxxx) from a seeded RNG.

Python >= 3.8.  openpyxl is only needed for the .xlsx file.
Usage: python scripts/build_iaa_v1.py [--seed 20261001] [--no-xlsx]
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import sys
from collections import Counter, OrderedDict
from pathlib import Path
from typing import Dict, List

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "annotation" / "iaa_v1"
MAIN = ROOT / "data/v3/instances_main.jsonl"
AUDIT = ROOT / "data/v3/instances_audit.jsonl"  # superset; holds author-rejected E1 rows dropped from main
T3 = ROOT / "data/t3/instances_inv_families.jsonl"
T7 = ROOT / "data/t7/instances_e1_unit.jsonl"
SEED = 20261001

TARGETS = OrderedDict(
    [
        ("conversion", 25),
        ("polarity_flip", 25),
        ("rephrase", 20),
        ("freq_route", 20),
        ("unit_edit", 20),
        ("value_edit", 20),
        ("base_claim", 20),
    ]
)

# Neutral description of what each edit is SUPPOSED to be (needed for Q1).
INTENDED = {
    "conversion": "UNIT CONVERSION - the dose was rewritten in another unit and should be the SAME amount (e.g. 0.5 g -> 500 mg).",
    "polarity_flip": "POLARITY FLIP - the edited claim should state the OPPOSITE meaning/recommendation, so the evidence now contradicts it.",
    "rephrase": "REPHRASE - wording changed; the meaning should be IDENTICAL to the original claim.",
    "freq_route": "FREQUENCY/ROUTE REWORDING - frequency or route written differently (e.g. 'twice daily' -> 'every 12 hours', 'intravenously' -> 'IV'); the meaning should be IDENTICAL.",
    "unit_edit_claim": "UNIT ERROR (in claim) - the unit was swapped with the number kept (e.g. 90 mg -> 90 g); the dose should now be WRONG and contradicted by the evidence.",
    "unit_edit_evid": "UNIT ERROR (in evidence) - the EVIDENCE unit was swapped with the number kept (e.g. 200 mg -> 200 g); claim unchanged; claim and evidence should now CONFLICT.",
    "value_edit": "DOSE VALUE ERROR - the dose number was multiplied (x2 or x10, e.g. 500 mg -> 1000 mg); the edited claim should now be clearly WRONG / contradicted by the evidence.",
    "base_claim": "BASE CLAIM (no edit) - judge the original claim itself: is it faithful to the evidence AND self-contained / well-specified (drug, population, dose clear enough to check)?",
}

CONDITION_LABEL = {
    "conversion": "Unit conversion",
    "polarity_flip": "Polarity flip",
    "rephrase": "Rephrase",
    "freq_route": "Frequency/route rewording",
    "unit_edit": "Unit error",
    "value_edit": "Dose value error",
    "base_claim": "Base claim",
}

Q1_TEXT = "Q1: Is the edited claim a valid instance of its intended condition? (Yes/No/Unsure)"
Q2_TEXT = "Q2 (conversions only): Is it clinically equivalent and natural? (Natural/Equivalent-unnatural/Not equivalent)"
Q3_TEXT = "Q3 (unit/value errors only): Clinical harm severity if this dose reached a patient (None/Minor/Moderate/Severe/Unsure)"
Q1_OPTS = ["Yes", "No", "Unsure"]
Q2_OPTS = ["Natural", "Equivalent-unnatural", "Not equivalent"]
Q3_OPTS = ["None", "Minor", "Moderate", "Severe", "Unsure"]
NA = "N/A"

SHEET_COLS = [
    ("item_id", "item_id"),
    ("condition", "condition"),
    ("intended_change", "intended_change"),
    ("evidence", "evidence"),
    ("original_evidence", "original_evidence (only if the evidence was edited)"),
    ("what_changed", "what_changed (automatic diff, for convenience)"),
    ("original_claim", "original_claim"),
    ("edited_claim", "edited_claim"),
    ("q1", Q1_TEXT),
    ("q2", Q2_TEXT),
    ("q3", Q3_TEXT),
    ("comment", "comment"),
]


def read_jsonl(path: Path) -> List[dict]:
    with path.open(encoding="utf-8") as fh:
        return [json.loads(l) for l in fh if l.strip()]


def read_csv(path: Path) -> List[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def yn(v: str) -> str:
    v = (v or "").strip().upper()
    return {"Y": "Yes", "N": "No"}.get(v, "")


def main110_bases() -> set:
    sys.path.insert(0, str(ROOT / "scripts"))
    try:
        from t15_card_report import load_main_dose_instances  # type: ignore

        return {r["base_id"] for r in load_main_dose_instances() if r.get("condition") == "C0"}
    except Exception as exc:  # pragma: no cover - fallback
        print("WARN: could not import t15_card_report (%s); using all main C0 DOSE bases" % exc)
        return {r["base_id"] for r in read_jsonl(MAIN) if r.get("type") == "DOSE" and r.get("condition") == "C0"}


def pick(rng: random.Random, pool: List[dict], k: int, key: str = "instance_id") -> List[dict]:
    pool = sorted(pool, key=lambda r: r[key])
    if k >= len(pool):
        return list(pool)
    return rng.sample(pool, k)


def build(seed: int) -> List[dict]:
    rng = random.Random(seed)
    main = read_jsonl(MAIN)
    by_id = {r["instance_id"]: r for r in main}
    c0 = {(r["base_id"]): r for r in main if r["condition"] == "C0"}
    b110 = main110_bases()
    items: List[dict] = []

    def add(stratum, sub, inst_id, base_id, magnitude, side, evidence, orig_ev, orig_claim, edited_claim,
            q1, q2, src, col, note, in110):
        items.append(
            dict(stratum=stratum, intended_key=sub, instance_id=inst_id, base_id=base_id, magnitude=magnitude,
                 edited_side=side, evidence=evidence, original_evidence=orig_ev, original_claim=orig_claim,
                 edited_claim=edited_claim, author_q1=q1, author_q2=q2, label_source=src, label_column=col,
                 author_note=note, in_main110=in110)
        )

    # 1) conversions: all natural (nat3=Y) + some unnatural (nat3=N) on the main 110 bases
    nat3 = {r["base_id"]: r for r in read_csv(ROOT / "qc/nat3_labels.csv")}
    conv = [r for r in main if r["condition"] == "INV-claim" and r.get("magnitude") == "unit_conversion"
            and r["base_id"] in b110 and r["base_id"] in nat3]
    conv_y = [r for r in conv if nat3[r["base_id"]]["natural"].strip().upper() == "Y"]
    conv_n = [r for r in conv if nat3[r["base_id"]]["natural"].strip().upper() == "N"]
    chosen = pick(rng, conv_y, len(conv_y)) + pick(rng, conv_n, TARGETS["conversion"] - len(conv_y))
    for r in chosen:
        lab = nat3[r["base_id"]]
        nat = lab["natural"].strip().upper()
        add("conversion", "conversion", r["instance_id"], r["base_id"], "unit_conversion", "claim",
            r["evidence"], "", r["claim_c0"], r["claim"], "Yes",
            "Natural" if nat == "Y" else "Equivalent-unnatural", "qc/nat3_labels.csv", "natural",
            lab.get("note", ""), True)

    # 2) polarity flips (author PI_ok on flips)
    flips = {r["base_id"]: r for r in read_csv(ROOT / "qc/e1_pol_review.csv")}
    flab = read_csv(ROOT / "qc/flip_PI_ok_labels.csv")
    fy = [dict(r, instance_id=r["base_id"]) for r in flab if r["PI_ok"].strip().upper() == "Y" and r["base_id"] in flips]
    fn = [dict(r, instance_id=r["base_id"]) for r in flab if r["PI_ok"].strip().upper() == "N" and r["base_id"] in flips]
    n_n = min(len(fn), TARGETS["polarity_flip"] // 2)
    chosen = pick(rng, fn, n_n) + pick(rng, fy, TARGETS["polarity_flip"] - n_n)
    for r in chosen:
        f = flips[r["base_id"]]
        base = c0[r["base_id"]]
        add("polarity_flip", "polarity_flip", r["base_id"] + "__flip", r["base_id"], r["flip_mechanism"], "claim",
            base["evidence"], "", f["original_claim"], f["flipped_claim"], yn(r["PI_ok"]), None,
            "qc/flip_PI_ok_labels.csv", "PI_ok", "", False)

    # 3) rephrases (POLARITY template rewrites, author PI_ok)
    reph = [dict(r, instance_id=r["base_id"]) for r in read_csv(ROOT / "qc/inv_manual_review.csv")
            if r["type"] == "POLARITY" and r["mechanism"] == "template" and yn(r["PI_ok"]) and r["base_id"] in c0]
    ry = [r for r in reph if yn(r["PI_ok"]) == "Yes"]
    rn = [r for r in reph if yn(r["PI_ok"]) == "No"]
    n_n = min(len(rn), TARGETS["rephrase"] // 2)
    chosen = pick(rng, rn, n_n) + pick(rng, ry, TARGETS["rephrase"] - n_n)
    for r in chosen:
        base = c0[r["base_id"]]
        add("rephrase", "rephrase", r["base_id"] + "__INV-claim__rephrase", r["base_id"], "rephrase", "claim",
            base["evidence"], "", r["original"], r["rewrite"], yn(r["PI_ok"]), None,
            "qc/inv_manual_review.csv", "PI_ok", "", False)

    # 4) frequency / route rewrites (author PI_ok per instance)
    t3 = {r["instance_id"]: r for r in read_jsonl(T3)}
    fr = [r for r in read_csv(ROOT / "qc/inv_claim_PI_ok_labels.csv")
          if r["family"] in ("inv_frequency", "inv_route") and r["instance_id"] in t3 and yn(r["PI_ok"])]
    half = TARGETS["freq_route"] // 2
    chosen = []
    for lab, k in (("No", half), ("Yes", TARGETS["freq_route"] - half)):
        fam_f = [r for r in fr if yn(r["PI_ok"]) == lab and r["family"] == "inv_frequency"]
        fam_r = [r for r in fr if yn(r["PI_ok"]) == lab and r["family"] == "inv_route"]
        kf = min(len(fam_f), k // 2)
        kr = min(len(fam_r), k - kf)
        kf = min(len(fam_f), k - kr)
        chosen += pick(rng, fam_f, kf) + pick(rng, fam_r, kr)
    for r in chosen:
        inst = t3[r["instance_id"]]
        add("freq_route", "freq_route", r["instance_id"], r["base_id"], r["family"], "claim",
            inst["evidence"], "", inst["claim_c0"], inst["claim"], yn(r["PI_ok"]), None,
            "qc/inv_claim_PI_ok_labels.csv", "PI_ok", "", r["base_id"] in b110)

    # 5) unit edits (T7 x1000 unit swaps; author PI_ok - all accepted in the source labels)
    t7 = {r["instance_id"]: r for r in read_jsonl(T7)}
    ulab = [r for r in read_csv(ROOT / "qc/t7_PI_ok_labels.csv") if r["instance_id"] in t7 and yn(r["PI_ok"])]
    for r in ulab:
        r["_base"] = t7[r["instance_id"]]["base_id"]
    u_claim = [r for r in ulab if "E1-unit-claim" in r["instance_id"] and r["_base"] in b110]
    u_evid = [r for r in ulab if "E1-unit-evid" in r["instance_id"] and r["_base"] in b110]
    used = set()
    chosen = pick(rng, u_claim, TARGETS["unit_edit"])
    used = {r["_base"] for r in chosen}
    rest = TARGETS["unit_edit"] - len(chosen)
    ev_new = [r for r in u_evid if r["_base"] not in used]
    ev_dup = [r for r in u_evid if r["_base"] in used]
    sel = pick(rng, ev_new, rest)
    if len(sel) < rest:
        sel += pick(rng, ev_dup, rest - len(sel))
    chosen += sel
    for r in chosen:
        inst = t7[r["instance_id"]]
        if "E1-unit-claim" in r["instance_id"]:
            add("unit_edit", "unit_edit_claim", r["instance_id"], inst["base_id"], inst.get("magnitude"), "claim",
                inst["evidence"], "", inst["claim_c0"], inst["claim"], yn(r["PI_ok"]), None,
                "qc/t7_PI_ok_labels.csv", "PI_ok", r.get("note", ""), True)
        else:
            add("unit_edit", "unit_edit_evid", r["instance_id"], inst["base_id"], inst.get("magnitude"), "evidence",
                inst["evidence"], inst["evidence_c0"], inst["claim_c0"], inst["claim"], yn(r["PI_ok"]), None,
                "qc/t7_PI_ok_labels.csv", "PI_ok", r.get("note", ""), True)

    # 6) value edits (E1-claim x2 / x10 on main bases; author dose3 E1x2_ok / E1x10_ok)
    d3 = {r["base_id"]: r for r in read_csv(ROOT / "qc/dose3_labels.csv")}
    # author-rejected E1 rows were removed from instances_main.jsonl; recover them from the audit superset
    audit = {r["instance_id"]: r for r in read_jsonl(AUDIT)} if AUDIT.is_file() else {}
    vlookup = dict(audit)
    vlookup.update(by_id)
    vpool = []
    for bid, lab in d3.items():
        if bid not in b110:
            continue
        for mag, col in (("x2", "E1x2_ok"), ("x10", "E1x10_ok")):
            iid = "%s__E1-claim__%s" % (bid, mag)
            if iid in vlookup and yn(lab[col]):
                vpool.append(dict(instance_id=iid, base_id=bid, magnitude=mag, col=col, lab=yn(lab[col]),
                                  note=lab.get("note", "")))
    vn = [r for r in vpool if r["lab"] == "No"]
    chosen = pick(rng, vn, min(len(vn), TARGETS["value_edit"] // 2))
    used = {r["base_id"] for r in chosen}
    need = TARGETS["value_edit"] - len(chosen)
    vy = [r for r in vpool if r["lab"] == "Yes"]
    # balance magnitudes and prefer bases not already shown as a value edit
    n_x2 = sum(r["magnitude"] == "x2" for r in chosen)
    target_x2 = TARGETS["value_edit"] // 2 - n_x2
    sel = []
    for mag, k in (("x2", max(0, target_x2)), ("x10", need - max(0, target_x2))):
        p_new = [r for r in vy if r["magnitude"] == mag and r["base_id"] not in used]
        s = pick(rng, p_new, k)
        used |= {r["base_id"] for r in s}
        sel += s
    if len(sel) < need:
        left = [r for r in vy if r not in sel]
        sel += pick(rng, left, need - len(sel))
    chosen += sel
    for r in chosen:
        inst = vlookup[r["instance_id"]]
        add("value_edit", "value_edit", r["instance_id"], r["base_id"], r["magnitude"], "claim",
            inst["evidence"], "", inst["claim_c0"], inst["claim"], r["lab"], None,
            "qc/dose3_labels.csv", r["col"], r["note"], True)

    # 7) base claims (C0) on main bases; author dose3 C0_ok
    cpool = [dict(instance_id="%s__C0__none" % b, base_id=b, lab=yn(l["C0_ok"]), note=l.get("note", ""))
             for b, l in d3.items() if b in b110 and yn(l["C0_ok"]) and b in c0]
    cy = [r for r in cpool if r["lab"] == "Yes"]
    cn = [r for r in cpool if r["lab"] == "No"]
    n_n = min(len(cn), TARGETS["base_claim"] // 2)
    chosen = pick(rng, cn, n_n) + pick(rng, cy, TARGETS["base_claim"] - n_n)
    for r in chosen:
        base = c0[r["base_id"]]
        add("base_claim", "base_claim", base["instance_id"], r["base_id"], "none", "none",
            base["evidence"], "", base["claim"], base["claim"], r["lab"], None,
            "qc/dose3_labels.csv", "C0_ok", r["note"], True)

    # shuffle + opaque ids
    rng.shuffle(items)
    ids = set()
    for it in items:
        while True:
            oid = "IAA1-%06x" % rng.randrange(16 ** 6)
            if oid not in ids:
                ids.add(oid)
                break
        it["item_id"] = oid
    return items


def word_diff(a: str, b: str) -> str:
    """Compact word-level diff 'old -> new' (pure presentation; carries no label)."""
    import difflib

    wa, wb = a.split(), b.split()
    parts = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=wa, b=wb, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        old = " ".join(wa[i1:i2]) or "(nothing)"
        new = " ".join(wb[j1:j2]) or "(nothing)"
        parts.append("[%s] -> [%s]" % (old, new))
    return "; ".join(parts) if parts else "(no visible change)"


def sheet_rows(items: List[dict]) -> List[Dict[str, str]]:
    rows = []
    for it in items:
        st = it["stratum"]
        rows.append(
            OrderedDict(
                item_id=it["item_id"],
                condition=CONDITION_LABEL[st],
                intended_change=INTENDED[it["intended_key"]],
                evidence=it["evidence"],
                original_evidence=it["original_evidence"] or "",
                what_changed=(
                    "" if st == "base_claim"
                    else ("EVIDENCE: " + word_diff(it["original_evidence"], it["evidence"]))
                    if it["edited_side"] == "evidence"
                    else word_diff(it["original_claim"], it["edited_claim"])
                ),
                original_claim=it["original_claim"],
                edited_claim=it["edited_claim"] if st != "base_claim" else "(no edit - judge the original claim)",
                q1="",
                q2="" if st == "conversion" else NA,
                q3="" if st in ("unit_edit", "value_edit") else NA,
                comment="",
            )
        )
    return rows


def write_csv(rows, path: Path):
    with path.open("w", newline="", encoding="utf-8-sig") as fh:  # BOM so Excel shows Vietnamese/Unicode correctly
        w = csv.writer(fh)
        w.writerow([h for _, h in SHEET_COLS])
        for r in rows:
            w.writerow([r[k] for k, _ in SHEET_COLS])


def write_xlsx(rows, path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    lines = [
        "IAA v1 annotation sheet - 'Same Dose, Different Verdict'",
        "",
        "Please read GUIDELINES_en.md (English) or GUIDELINES_vi.md (Tieng Viet) before starting. Estimated time: 2-3 hours.",
        "Work in the 'Annotation' sheet. Fill one row per item; use the dropdowns in the Q1/Q2/Q3 columns.",
        "",
        Q1_TEXT + "  -> every row.",
        Q2_TEXT + "  -> only rows with condition 'Unit conversion' (others pre-filled N/A).",
        Q3_TEXT + "  -> only rows with condition 'Unit error' or 'Dose value error' (others pre-filled N/A).",
        "",
        "Severity (Q3, NCC MERP-inspired): None = no harm expected; Minor = could cause mild, temporary harm or need extra monitoring;",
        "Moderate = temporary harm needing treatment/intervention or prolonged stay; Severe = permanent harm, life-threatening, or death; Unsure.",
        "",
        "Judge each item independently. Do not search for the source of the claim. Use 'comment' for anything unclear.",
        "Do not reorder, delete or rename rows/columns. Save as .xlsx (or .csv) and return the file.",
    ]
    for i, l in enumerate(lines, 1):
        ins.cell(row=i, column=1, value=l)
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 150

    ws = wb.create_sheet("Annotation")
    widths = {"item_id": 13, "condition": 16, "intended_change": 38, "evidence": 70, "original_evidence": 45, "what_changed": 30,
              "original_claim": 50, "edited_claim": 50, "q1": 16, "q2": 22, "q3": 16, "comment": 40}
    hdr_fill = PatternFill("solid", fgColor="DDEBF7")
    na_fill = PatternFill("solid", fgColor="EEEEEE")
    ans_fill = PatternFill("solid", fgColor="FFF2CC")
    for j, (k, h) in enumerate(SHEET_COLS, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = widths[k]
    ws.row_dimensions[1].height = 75
    dv1 = DataValidation(type="list", formula1='"%s"' % ",".join(Q1_OPTS), allow_blank=True, showDropDown=False)
    dv2 = DataValidation(type="list", formula1='"%s"' % ",".join(Q2_OPTS + [NA]), allow_blank=True)
    dv3 = DataValidation(type="list", formula1='"%s"' % ",".join(Q3_OPTS + [NA]), allow_blank=True)
    for dv, name in ((dv1, "Q1"), (dv2, "Q2"), (dv3, "Q3")):
        dv.error = "Please choose a value from the dropdown list."
        dv.errorTitle = "Invalid %s value" % name
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
    col_idx = {k: j for j, (k, _) in enumerate(SHEET_COLS, 1)}
    for i, r in enumerate(rows, 2):
        for k, j in col_idx.items():
            c = ws.cell(row=i, column=j, value=r[k])
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for k, dv in (("q1", dv1), ("q2", dv2), ("q3", dv3)):
            cell = ws.cell(row=i, column=col_idx[k])
            if r[k] == NA:
                cell.fill = na_fill
            else:
                cell.fill = ans_fill
                dv.add(cell.coordinate)
        ws.cell(row=i, column=col_idx["comment"]).fill = ans_fill
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:%s%d" % (ws.cell(row=1, column=len(SHEET_COLS)).column_letter, len(rows) + 1)
    wb.active = 1
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--seed", type=int, default=SEED)
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--out-dir", default=str(OUT_DIR))
    args = ap.parse_args()
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    items = build(args.seed)
    rows = sheet_rows(items)
    write_csv(rows, out / "annotation_sheet.csv")
    if not args.no_xlsx:
        write_xlsx(rows, out / "annotation_sheet.xlsx")
    key_fields = ["item_id", "stratum", "intended_key", "instance_id", "base_id", "magnitude", "edited_side",
                  "author_q1", "author_q2", "author_q3", "label_source", "label_column", "author_note", "in_main110"]
    with (out / "private_key.jsonl").open("w", encoding="utf-8") as fh:
        for it in sorted(items, key=lambda x: x["item_id"]):
            it.setdefault("author_q3", None)
            fh.write(json.dumps({k: it.get(k) for k in key_fields}, ensure_ascii=False) + "\n")
    counts = Counter(it["stratum"] for it in items)
    acc = Counter((it["stratum"], it["author_q1"]) for it in items)
    q2 = Counter(it["author_q2"] for it in items if it["stratum"] == "conversion")
    manifest = OrderedDict(
        version="iaa_v1",
        seed=args.seed,
        n_items=len(items),
        targets=TARGETS,
        counts_per_stratum=OrderedDict((s, counts.get(s, 0)) for s in TARGETS),
        note=("Accepted/rejected balance per stratum is recorded only in private_key.jsonl. "
              "Sheet is blind: no author labels, no rule outputs, no instance ids, shuffled."),
        sources=OrderedDict(
            conversion="data/v3/instances_main.jsonl INV-claim unit_conversion (main 110) + qc/nat3_labels.csv",
            polarity_flip="qc/e1_pol_review.csv text + qc/flip_PI_ok_labels.csv; evidence = POLARITY C0",
            rephrase="qc/inv_manual_review.csv (POLARITY/template); evidence = POLARITY C0",
            freq_route="data/t3/instances_inv_families.jsonl + qc/inv_claim_PI_ok_labels.csv (inv_frequency, inv_route)",
            unit_edit="data/t7/instances_e1_unit.jsonl + qc/t7_PI_ok_labels.csv (main 110 bases)",
            value_edit="data/v3/instances_main.jsonl E1-claim x2/x10 + qc/dose3_labels.csv E1x2_ok/E1x10_ok",
            base_claim="data/v3/instances_main.jsonl C0 + qc/dose3_labels.csv C0_ok",
        ),
        files=["annotation_sheet.xlsx", "annotation_sheet.csv", "GUIDELINES_en.md", "GUIDELINES_vi.md",
               "private_key.jsonl (DO NOT SHARE with annotators)"],
    )
    (out / "sample_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("n_items", len(items))
    for s in TARGETS:
        print("  %-14s %3d  author Yes=%d No=%d" % (s, counts.get(s, 0), acc.get((s, "Yes"), 0), acc.get((s, "No"), 0)))
    print("  conversion author_q2:", dict(q2))


if __name__ == "__main__":
    main()
