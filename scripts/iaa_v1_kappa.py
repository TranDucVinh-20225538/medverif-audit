#!/usr/bin/env python3
"""IAA v1 agreement: filled annotation sheet vs the author's private key.

Reads annotation/iaa_v1/annotation_sheet.{xlsx,csv} as filled by an annotator
and annotation/iaa_v1/private_key.jsonl (opaque item id -> instance id +
author labels) and reports:

  * Q1 (valid instance of intended condition; nominal Yes/No):
      Cohen's kappa per stratum and overall with percentile bootstrap CIs,
      raw agreement, Gwet's AC1 (prevalence-robust; kappa is undefined or
      unstable when the author labels in a stratum are all one class),
      confusion matrices (author rows x annotator cols, incl. Unsure).
      Primary analysis drops Unsure; a sensitivity analysis maps Unsure -> No.
  * Q2 (conversions; ordinal Natural > Equivalent-unnatural > Not equivalent):
      unweighted + linear-weighted kappa vs author nat3 labels, CI, confusion.
  * Q3 (unit vs value errors; ordinal None < Minor < Moderate < Severe):
      severity distributions per stratum and magnitude; quadratic-weighted
      kappa if author severities exist in the key (author_q3) and/or if a
      second annotator sheet is given with --sheet2.
  * With --sheet2: annotator-vs-annotator kappa for Q1/Q2/Q3 as well.

Python >= 3.8, stdlib only (openpyxl only to read .xlsx).

  python scripts/iaa_v1_kappa.py --sheet filled.xlsx \
      [--key annotation/iaa_v1/private_key.jsonl] [--sheet2 other.xlsx] \
      [--out-md notes/iaa_v1_report.md] [--out-json notes/iaa_v1_report.json]
  python scripts/iaa_v1_kappa.py --simulate OUT.xlsx --agree 0.85   # test fixture
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
from collections import Counter, OrderedDict
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

ROOT = Path(__file__).resolve().parents[1]
KEY = ROOT / "annotation/iaa_v1/private_key.jsonl"
SHEET = ROOT / "annotation/iaa_v1/annotation_sheet.xlsx"
STRATA = ["conversion", "polarity_flip", "rephrase", "freq_route", "unit_edit", "value_edit", "base_claim"]
Q1C = ["Yes", "No"]
Q2C = ["Natural", "Equivalent-unnatural", "Not equivalent"]
Q3C = ["None", "Minor", "Moderate", "Severe"]
NA = "N/A"


# ----------------------------------------------------------------- io
def _norm(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    low = s.lower()
    canon = {"yes": "Yes", "y": "Yes", "no": "No", "n": "No", "unsure": "Unsure", "n/a": NA, "na": NA,
             "natural": "Natural", "equivalent-unnatural": "Equivalent-unnatural",
             "equivalent - unnatural": "Equivalent-unnatural", "not equivalent": "Not equivalent",
             "none": "None", "minor": "Minor", "moderate": "Moderate", "severe": "Severe"}
    return canon.get(low, s)


def _colmap(header: Sequence[str]) -> Dict[str, int]:
    out = {}
    for i, h in enumerate(header):
        h0 = (h or "").strip().lower()
        if h0.startswith("item_id"):
            out["item_id"] = i
        elif h0.startswith("q1"):
            out["q1"] = i
        elif h0.startswith("q2"):
            out["q2"] = i
        elif h0.startswith("q3"):
            out["q3"] = i
        elif h0.startswith("comment"):
            out["comment"] = i
    missing = {"item_id", "q1", "q2", "q3"} - set(out)
    if missing:
        raise SystemExit("sheet is missing columns: %s" % sorted(missing))
    return out


def read_sheet(path: Path) -> Dict[str, Dict[str, str]]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Annotation"] if "Annotation" in wb.sheetnames else wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.reader(fh))
    cm = _colmap([str(x) if x is not None else "" for x in rows[0]])
    out = {}
    for r in rows[1:]:
        if not r or r[cm["item_id"]] in (None, ""):
            continue
        get = lambda k: _norm(r[cm[k]]) if cm.get(k) is not None and cm[k] < len(r) else ""
        out[str(r[cm["item_id"]]).strip()] = {"q1": get("q1"), "q2": get("q2"), "q3": get("q3"),
                                               "comment": get("comment")}
    return out


def read_key(path: Path) -> Dict[str, dict]:
    with Path(path).open(encoding="utf-8") as fh:
        return {d["item_id"]: d for d in (json.loads(l) for l in fh if l.strip())}


# ----------------------------------------------------------------- stats
def cohen_kappa(a: Sequence[str], b: Sequence[str], cats: Sequence[str], weights: str = "none") -> float:
    """Cohen's kappa; weights in {'none','linear','quadratic'} over the order of cats."""
    n = len(a)
    if n == 0:
        return float("nan")
    k = len(cats)
    idx = {c: i for i, c in enumerate(cats)}
    obs = [[0.0] * k for _ in range(k)]
    for x, y in zip(a, b):
        obs[idx[x]][idx[y]] += 1
    ra = [sum(obs[i]) / n for i in range(k)]
    cb = [sum(obs[i][j] for i in range(k)) / n for j in range(k)]

    def w(i, j):
        if weights == "none" or k == 1:
            return 0.0 if i == j else 1.0
        d = abs(i - j) / (k - 1)
        return d if weights == "linear" else d * d

    do = sum(w(i, j) * obs[i][j] / n for i in range(k) for j in range(k))
    de = sum(w(i, j) * ra[i] * cb[j] for i in range(k) for j in range(k))
    if de == 0:
        return float("nan")  # undefined: no expected disagreement (e.g. both raters constant)
    return 1.0 - do / de


def gwet_ac1(a: Sequence[str], b: Sequence[str], cats: Sequence[str]) -> float:
    n, q = len(a), len(cats)
    if n == 0 or q < 2:
        return float("nan")
    pa = sum(x == y for x, y in zip(a, b)) / n
    pi = [(sum(x == c for x in a) + sum(y == c for y in b)) / (2.0 * n) for c in cats]
    pe = sum(p * (1 - p) for p in pi) / (q - 1)
    return float("nan") if pe == 1 else (pa - pe) / (1 - pe)


def raw_agree(a, b) -> float:
    return sum(x == y for x, y in zip(a, b)) / len(a) if a else float("nan")


def bootstrap(fn, a, b, n_boot: int, seed: int) -> Tuple[float, float, int]:
    rng = random.Random(seed)
    n = len(a)
    vals = []
    for _ in range(n_boot):
        ix = [rng.randrange(n) for _ in range(n)]
        v = fn([a[i] for i in ix], [b[i] for i in ix])
        if not math.isnan(v):
            vals.append(v)
    if len(vals) < max(20, n_boot // 10):
        return float("nan"), float("nan"), len(vals)
    vals.sort()
    lo = vals[int(0.025 * (len(vals) - 1))]
    hi = vals[int(math.ceil(0.975 * (len(vals) - 1)))]
    return lo, hi, len(vals)


def confusion(a, b, rows, cols) -> List[List[int]]:
    c = Counter(zip(a, b))
    return [[c.get((r, k), 0) for k in cols] for r in rows]


def fmt(x, nd=3) -> str:
    return "n/a" if x is None or (isinstance(x, float) and math.isnan(x)) else ("%.*f" % (nd, x))


def md_conf(title, m, rows, cols, rlab="author", clab="annotator") -> List[str]:
    out = ["", "_%s_ (rows = %s, cols = %s)" % (title, rlab, clab), "",
           "| | " + " | ".join(cols) + " |", "|---|" + "---|" * len(cols)]
    for r, line in zip(rows, m):
        out.append("| **%s** | " % r + " | ".join(str(v) for v in line) + " |")
    return out


# ----------------------------------------------------------------- analysis
def pair_stats(a, b, cats, weights, n_boot, seed, ac1=False) -> dict:
    kfn = lambda x, y: cohen_kappa(x, y, cats, weights)
    lo, hi, nb = bootstrap(kfn, a, b, n_boot, seed) if a else (float("nan"), float("nan"), 0)
    d = OrderedDict(n=len(a), raw_agreement=raw_agree(a, b), kappa=kfn(a, b) if a else float("nan"),
                    ci95=[lo, hi], boot_valid=nb, weights=weights)
    if weights != "none":
        d["kappa_unweighted"] = cohen_kappa(a, b, cats, "none") if a else float("nan")
    if ac1:
        d["gwet_ac1"] = gwet_ac1(a, b, cats)
        lo2, hi2, _ = bootstrap(lambda x, y: gwet_ac1(x, y, cats), a, b, n_boot, seed + 1) if a else (float("nan"),) * 3
        d["gwet_ac1_ci95"] = [lo2, hi2]
    return d


def analyse(key, sheet, sheet2=None, n_boot=2000, seed=20261001) -> Tuple[dict, List[str]]:
    res = OrderedDict()
    md = ["# IAA v1 agreement report", "",
          "Bootstrap: %d percentile resamples over items, seed %d. Kappa = n/a when undefined "
          "(no expected disagreement, e.g. all author labels in the stratum are one class); "
          "Gwet's AC1 is reported alongside because it is robust to such prevalence skew." % (n_boot, seed), ""]
    ids = sorted(set(key) & set(sheet))
    miss = sorted(set(key) - set(sheet))
    extra = sorted(set(sheet) - set(key))
    blank_q1 = [i for i in ids if sheet[i]["q1"] not in ("Yes", "No", "Unsure")]
    res["coverage"] = dict(key_items=len(key), sheet_items=len(sheet), matched=len(ids), missing=len(miss),
                           unknown_ids=len(extra), q1_blank_or_invalid=len(blank_q1))
    md += ["## Coverage", "", "key items %d; sheet rows %d; matched %d; missing from sheet %d; unknown ids %d; "
           "Q1 blank/invalid %d" % (len(key), len(sheet), len(ids), len(miss), len(extra), len(blank_q1)), ""]

    # ---- Q1
    md += ["## Q1 - valid instance of the intended condition (author vs annotator)", "",
           "Primary: Unsure excluded. Sensitivity: Unsure -> No.", "",
           "| stratum | n | Unsure | author Yes/No | raw agr. | kappa [95% CI] | AC1 [95% CI] | kappa (Unsure->No) |",
           "|---|---|---|---|---|---|---|---|"]
    q1 = OrderedDict()
    confs = []
    for st in STRATA + ["ALL"]:
        sel = [i for i in ids if (st == "ALL" or key[i]["stratum"] == st) and key[i].get("author_q1") in Q1C
               and sheet[i]["q1"] in ("Yes", "No", "Unsure")]
        a_all = [key[i]["author_q1"] for i in sel]
        b_all = [sheet[i]["q1"] for i in sel]
        a = [x for x, y in zip(a_all, b_all) if y != "Unsure"]
        b = [y for y in b_all if y != "Unsure"]
        d = pair_stats(a, b, Q1C, "none", n_boot, seed, ac1=True)
        d["n_unsure"] = b_all.count("Unsure")
        d["author_counts"] = dict(Counter(a_all))
        d["kappa_unsure_as_no"] = cohen_kappa(a_all, ["No" if y == "Unsure" else y for y in b_all], Q1C) if a_all else float("nan")
        d["confusion"] = confusion(a_all, b_all, Q1C, Q1C + ["Unsure"])
        if len(set(a_all)) < 2:
            # author labels are all one class: kappa is mathematically 0 or undefined and not informative
            d["kappa"] = d["kappa_unsure_as_no"] = float("nan")
            d["ci95"] = [float("nan"), float("nan")]
            d["note"] = "author labels single-class; kappa not informative - use raw agreement / AC1"
        q1[st] = d
        md.append("| %s | %d | %d | %d/%d | %s | %s [%s, %s] | %s [%s, %s] | %s |" % (
            st, len(a_all), d["n_unsure"], d["author_counts"].get("Yes", 0), d["author_counts"].get("No", 0),
            fmt(d["raw_agreement"]), fmt(d["kappa"]), fmt(d["ci95"][0]), fmt(d["ci95"][1]),
            fmt(d["gwet_ac1"]), fmt(d["gwet_ac1_ci95"][0]), fmt(d["gwet_ac1_ci95"][1]), fmt(d["kappa_unsure_as_no"])))
        confs += md_conf("Q1 confusion, %s" % st, d["confusion"], Q1C, Q1C + ["Unsure"])
    res["q1"] = q1
    single = [st for st in STRATA if q1[st].get("note")]
    if single:
        md += ["", "Strata with single-class author labels (kappa shown as n/a): %s." % ", ".join(single)]
    md += ["", "### Q1 confusion matrices"] + confs + [""]

    # ---- Q2
    sel = [i for i in ids if key[i]["stratum"] == "conversion" and key[i].get("author_q2") in Q2C
           and sheet[i]["q2"] in Q2C]
    a = [key[i]["author_q2"] for i in sel]
    b = [sheet[i]["q2"] for i in sel]
    d = pair_stats(a, b, Q2C, "linear", n_boot, seed, ac1=True)
    d["confusion"] = confusion(a, b, Q2C, Q2C)
    d["annotator_counts"] = dict(Counter(b))
    d["n_blank_or_invalid"] = sum(1 for i in ids if key[i]["stratum"] == "conversion" and sheet[i]["q2"] not in Q2C)
    res["q2"] = d
    md += ["## Q2 - conversions: equivalent and natural (ordinal, linear-weighted kappa)", "",
           "n %d; raw agreement %s; weighted kappa %s [%s, %s]; unweighted kappa %s; AC1 %s; blank/invalid %d" % (
               d["n"], fmt(d["raw_agreement"]), fmt(d["kappa"]), fmt(d["ci95"][0]), fmt(d["ci95"][1]),
               fmt(d["kappa_unweighted"]), fmt(d["gwet_ac1"]), d["n_blank_or_invalid"]),
           "", "Author labels exist only as Natural (nat3 = Y) vs Equivalent-unnatural (nat3 = N); "
           "'Not equivalent' from the annotator flags a conversion error."]
    md += md_conf("Q2 confusion", d["confusion"], Q2C, Q2C) + [""]

    # ---- Q3
    md += ["## Q3 - harm severity (unit vs value errors)", ""]
    q3 = OrderedDict()
    cats_all = Q3C + ["Unsure"]
    md += ["| group | n | " + " | ".join(cats_all) + " | mean ordinal (0-3, excl. Unsure) |", "|---|---|" + "---|" * (len(cats_all) + 1)]
    groups = [("unit_edit", None), ("value_edit", None), ("value_edit", "x2"), ("value_edit", "x10"),
              ("value_edit|author_q1=Yes", None), ("value_edit|author_q1=No", None)]
    for g, mag in groups:
        st = g.split("|")[0]
        sel = [i for i in ids if key[i]["stratum"] == st and (mag is None or key[i]["magnitude"] == mag)]
        if "|" in g:
            sel = [i for i in sel if key[i].get("author_q1") == g.split("=")[1]]
        vals = [sheet[i]["q3"] for i in sel if sheet[i]["q3"] in cats_all]
        c = Counter(vals)
        ords = [Q3C.index(v) for v in vals if v in Q3C]
        name = g + ("" if mag is None else " " + mag)
        q3[name] = dict(n=len(vals), counts={k: c.get(k, 0) for k in cats_all},
                        mean_ordinal=(sum(ords) / len(ords)) if ords else float("nan"))
        md.append("| %s | %d | %s | %s |" % (name, len(vals), " | ".join(str(c.get(k, 0)) for k in cats_all),
                                            fmt(q3[name]["mean_ordinal"], 2)))
    # unit vs value difference in mean ordinal severity, bootstrap CI
    u = [Q3C.index(sheet[i]["q3"]) for i in ids if key[i]["stratum"] == "unit_edit" and sheet[i]["q3"] in Q3C]
    v = [Q3C.index(sheet[i]["q3"]) for i in ids if key[i]["stratum"] == "value_edit" and sheet[i]["q3"] in Q3C]
    if u and v:
        rng = random.Random(seed + 7)
        diffs = sorted(sum(rng.choice(u) for _ in u) / len(u) - sum(rng.choice(v) for _ in v) / len(v) for _ in range(n_boot))
        q3["unit_minus_value_mean"] = dict(diff=sum(u) / len(u) - sum(v) / len(v),
                                           ci95=[diffs[int(0.025 * (n_boot - 1))], diffs[int(math.ceil(0.975 * (n_boot - 1)))]])
        md += ["", "Mean severity unit - value = %s [95%% CI %s, %s]" % (
            fmt(q3["unit_minus_value_mean"]["diff"], 2), fmt(q3["unit_minus_value_mean"]["ci95"][0], 2),
            fmt(q3["unit_minus_value_mean"]["ci95"][1], 2))]
    sev_ids = [i for i in ids if key[i]["stratum"] in ("unit_edit", "value_edit")]
    with_author = [i for i in sev_ids if key[i].get("author_q3") in Q3C and sheet[i]["q3"] in Q3C]
    if with_author:
        d = pair_stats([key[i]["author_q3"] for i in with_author], [sheet[i]["q3"] for i in with_author],
                       Q3C, "quadratic", n_boot, seed)
        q3["vs_author_weighted_kappa"] = d
        md += ["", "Q3 vs author (quadratic-weighted, Unsure excluded): n %d, kappa_w %s [%s, %s], raw %s" % (
            d["n"], fmt(d["kappa"]), fmt(d["ci95"][0]), fmt(d["ci95"][1]), fmt(d["raw_agreement"]))]
    else:
        md += ["", "No author severity labels in the key (author_q3 is null) -> Q3 reported as distributions only; "
               "weighted kappa is computed when a second sheet (--sheet2) or author_q3 values are provided."]
    res["q3"] = q3

    # ---- annotator vs annotator
    if sheet2 is not None:
        md += ["", "## Annotator 1 vs annotator 2", "", "| question | n | raw | kappa [95% CI] | weighting |", "|---|---|---|---|---|"]
        aa = OrderedDict()
        for q, cats, w, strata in (("q1", Q1C, "none", None), ("q2", Q2C, "linear", ("conversion",)),
                                   ("q3", Q3C, "quadratic", ("unit_edit", "value_edit"))):
            sel = [i for i in ids if i in sheet2 and (strata is None or key[i]["stratum"] in strata)
                   and sheet[i][q] in cats and sheet2[i][q] in cats]
            d = pair_stats([sheet[i][q] for i in sel], [sheet2[i][q] for i in sel], cats, w, n_boot, seed)
            aa[q] = d
            md.append("| %s | %d | %s | %s [%s, %s] | %s |" % (q, d["n"], fmt(d["raw_agreement"]), fmt(d["kappa"]),
                                                             fmt(d["ci95"][0]), fmt(d["ci95"][1]), w))
        res["annotator_vs_annotator"] = aa
    return res, md


# ----------------------------------------------------------------- simulation (test fixture)
def simulate(key: Dict[str, dict], template: Path, out: Path, agree: float, seed: int) -> None:
    rng = random.Random(seed)
    ans = {}
    for iid, k in key.items():
        st = k["stratum"]
        a1 = k.get("author_q1") or "Yes"
        r = rng.random()
        q1 = a1 if r < agree else ("Unsure" if r < agree + (1 - agree) / 3 else ("No" if a1 == "Yes" else "Yes"))
        q2 = NA
        if st == "conversion":
            a2 = k.get("author_q2") or "Natural"
            q2 = a2 if rng.random() < agree else rng.choice([c for c in Q2C if c != a2])
        q3 = NA
        if st in ("unit_edit", "value_edit"):
            base = {"x1000": [0.05, 0.1, 0.25, 0.55, 0.05], "x10": [0.1, 0.2, 0.35, 0.3, 0.05],
                    "x2": [0.3, 0.35, 0.25, 0.05, 0.05]}.get(k.get("magnitude"), [0.2] * 5)
            q3 = rng.choices(Q3C + ["Unsure"], weights=base)[0]
        ans[iid] = (q1, q2, q3, "simulated")
    if template.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(template)
        ws = wb["Annotation"]
        header = [c.value for c in ws[1]]
        cm = _colmap([str(h) for h in header])
        for row in ws.iter_rows(min_row=2):
            iid = row[cm["item_id"]].value
            if iid in ans:
                for k, v in zip(("q1", "q2", "q3", "comment"), ans[iid]):
                    row[cm[k]].value = v
        wb.save(out)
    else:
        with template.open(newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.reader(fh))
        cm = _colmap(rows[0])
        for r in rows[1:]:
            if r[cm["item_id"]] in ans:
                for k, v in zip(("q1", "q2", "q3", "comment"), ans[r[cm["item_id"]]]):
                    r[cm[k]] = v
        with out.open("w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerows(rows)


def _jsonable(o):
    if isinstance(o, float) and math.isnan(o):
        return None
    if isinstance(o, dict):
        return {k: _jsonable(v) for k, v in o.items()}
    if isinstance(o, list):
        return [_jsonable(v) for v in o]
    return o


def main() -> None:
    ap = argparse.ArgumentParser(description="IAA v1 agreement (Cohen's kappa etc.)")
    ap.add_argument("--sheet", help="filled annotation sheet (.xlsx or .csv)")
    ap.add_argument("--sheet2", help="optional second annotator sheet")
    ap.add_argument("--key", default=str(KEY))
    ap.add_argument("--boot", type=int, default=2000)
    ap.add_argument("--seed", type=int, default=20261001)
    ap.add_argument("--out-md")
    ap.add_argument("--out-json")
    ap.add_argument("--simulate", help="write a simulated filled sheet to this path and exit")
    ap.add_argument("--template", default=str(SHEET))
    ap.add_argument("--agree", type=float, default=0.85)
    args = ap.parse_args()
    key = read_key(Path(args.key))
    if args.simulate:
        simulate(key, Path(args.template), Path(args.simulate), args.agree, args.seed)
        print("wrote simulated sheet", args.simulate)
        return
    if not args.sheet:
        raise SystemExit("--sheet is required")
    sheet = read_sheet(Path(args.sheet))
    sheet2 = read_sheet(Path(args.sheet2)) if args.sheet2 else None
    res, md = analyse(key, sheet, sheet2, args.boot, args.seed)
    md_text = "\n".join(md) + "\n"
    if args.out_md:
        Path(args.out_md).write_text(md_text, encoding="utf-8")
    if args.out_json:
        Path(args.out_json).write_text(json.dumps(_jsonable(res), indent=2) + "\n", encoding="utf-8")
    print(md_text)


if __name__ == "__main__":
    main()
